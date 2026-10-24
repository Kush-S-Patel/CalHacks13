"""
Integrated spreadsheet analyzer.

Maps messy account / revenue / AP / AR / supply-chain sheets into ERP entities.
Column detection is heuristic — works on real export messiness, not just demos.
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from erp.models import GLAccount, JournalLine, MoneyFlow, Party, new_id

# Header synonyms → canonical field
HEADER_MAP: dict[str, tuple[str, ...]] = {
    "account_code": ("account", "acct", "gl", "gl_code", "account_code", "acct_no", "number", "code"),
    "account_name": ("account_name", "name", "description", "acct_name", "title"),
    "account_type": ("type", "account_type", "category", "class", "classification"),
    "balance": ("balance", "amount", "ending_balance", "balance_usd", "total", "value"),
    "debit": ("debit", "dr", "debit_amount"),
    "credit": ("credit", "cr", "credit_amount"),
    "vendor": ("vendor", "supplier", "payee", "merchant"),
    "customer": ("customer", "client", "buyer", "account_name"),
    "revenue": ("revenue", "sales", "income", "net_sales", "arr"),
    "tier": ("tier", "level", "stage", "echelon"),
    "from_party": ("from", "source", "origin", "ship_from", "upstream"),
    "to_party": ("to", "target", "destination", "ship_to", "downstream"),
    "role": ("role", "party_type", "entity_type"),
    "memo": ("memo", "notes", "narrative", "description"),
    "date": ("date", "txn_date", "posting_date", "period"),
    "invoice": ("invoice", "invoice_id", "inv", "bill"),
    "fraud": ("fraud", "flag", "anomaly", "risk_flag"),
}


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", h.strip().lower()).strip("_")


def map_headers(headers: list[str]) -> dict[str, int]:
    """Return canonical_field -> column index."""
    out: dict[str, int] = {}
    normalized = [_norm(h) for h in headers]
    for canon, aliases in HEADER_MAP.items():
        for i, h in enumerate(normalized):
            if h in aliases or any(a in h for a in aliases):
                out.setdefault(canon, i)
                break
    return out


def parse_money(raw: Any) -> int:
    """Parse currency-ish cell to integer cents."""
    if raw is None or raw == "":
        return 0
    if isinstance(raw, (int, float)):
        # Heuristic: large ints without decimals already cents? Prefer dollars.
        v = float(raw)
        if abs(v) > 1_000_000 and float(raw) == int(raw) and abs(int(raw)) > 100_000:
            # likely already cents from export
            return int(raw)
        return int(round(v * 100))
    s = str(raw).strip()
    neg = s.startswith("(") and s.endswith(")")
    s = s.replace("$", "").replace(",", "").replace("(", "").replace(")", "").replace("USD", "").strip()
    if not s or s == "-":
        return 0
    try:
        v = float(s)
    except ValueError:
        return 0
    cents = int(round(v * 100))
    return -cents if neg else cents


def detect_kind(filename: str, headers: list[str]) -> str:
    fn = filename.lower()
    joined = " ".join(_norm(h) for h in headers)
    if any(k in fn for k in ("coa", "chart", "accounts", "ledger")):
        return "chart_of_accounts"
    if any(k in fn for k in ("revenue", "sales", "pnl", "p_and_l", "income")):
        return "revenue"
    if "ap" in fn or "payable" in fn or "aging_ap" in fn:
        return "ap_aging"
    if "ar" in fn or "receivable" in fn or "aging_ar" in fn:
        return "ar_aging"
    if "supply" in fn or "chain" in fn or "vendor_map" in fn:
        return "supply_chain"
    if "bank" in fn or "statement" in fn:
        return "bank"
    if "scan" in fn or "ocr" in fn or "receipt" in fn:
        return "scanned_financial"
    if "account" in joined and ("balance" in joined or "type" in joined):
        return "chart_of_accounts"
    if "vendor" in joined and "balance" in joined:
        return "ap_aging"
    if "customer" in joined:
        return "ar_aging"
    if "from" in joined and "to" in joined:
        return "supply_chain"
    if "revenue" in joined or "sales" in joined:
        return "revenue"
    return "other"


def slug_id(prefix: str, name: str) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:40] or "party"
    return f"{prefix}-{base}"


def analyze_rows(
    *,
    filename: str,
    headers: list[str],
    rows: list[list[Any]],
    kind_hint: str | None = None,
) -> dict[str, Any]:
    kind = kind_hint or detect_kind(filename, headers)
    cols = map_headers(headers)
    accounts: list[GLAccount] = []
    parties: list[Party] = []
    flows: list[MoneyFlow] = []
    journals: list[JournalLine] = []
    notes: list[str] = []

    def cell(row: list[Any], key: str, default: Any = "") -> Any:
        idx = cols.get(key)
        if idx is None or idx >= len(row):
            return default
        return row[idx]

    if kind == "chart_of_accounts" or ("account_code" in cols and "balance" in cols):
        kind = "chart_of_accounts"
        for row in rows:
            code = str(cell(row, "account_code") or "").strip()
            if not code:
                continue
            name = str(cell(row, "account_name") or code).strip()
            atype = str(cell(row, "account_type") or "asset").strip().lower()
            if "rev" in atype or "income" in atype or "sales" in atype:
                atype = "revenue"
            elif "exp" in atype or "cogs" in atype or "cost" in atype:
                atype = "expense"
            elif "liab" in atype:
                atype = "liability"
            elif "equity" in atype:
                atype = "equity"
            else:
                atype = "asset" if atype not in {"asset", "liability", "equity", "revenue", "expense"} else atype
            bal = parse_money(cell(row, "balance"))
            if not bal and ("debit" in cols or "credit" in cols):
                bal = parse_money(cell(row, "debit")) - parse_money(cell(row, "credit"))
            accounts.append(GLAccount(code=code, name=name, account_type=atype, balance_cents=bal))
            journals.append(
                JournalLine(
                    id=new_id("je"),
                    account_code=code,
                    debit_cents=max(bal, 0),
                    credit_cents=max(-bal, 0),
                    memo=f"Imported from {filename}",
                )
            )
        notes.append(f"Mapped {len(accounts)} GL accounts")

    elif kind in ("ap_aging", "ar_aging") or "vendor" in cols or "customer" in cols:
        role = "vendor" if kind == "ap_aging" or "vendor" in cols else "customer"
        name_key = "vendor" if role == "vendor" else "customer"
        if name_key not in cols and "account_name" in cols:
            name_key = "account_name"
        for row in rows:
            name = str(cell(row, name_key) or "").strip()
            if not name:
                continue
            amt = parse_money(cell(row, "balance") or cell(row, "revenue"))
            pid = slug_id("vnd" if role == "vendor" else "cust", name)
            parties.append(
                Party(
                    id=pid,
                    name=name,
                    role=role,
                    tier=2 if role == "vendor" else 6,
                    amount_cents=abs(amt),
                    meta={"source_file": filename},
                )
            )
            hub = "maya-bakery"
            if role == "vendor":
                flows.append(
                    MoneyFlow(
                        id=new_id("flow"),
                        source_id=hub,
                        target_id=pid,
                        amount_cents=abs(amt),
                        kind="payable",
                        memo=str(cell(row, "memo") or "AP"),
                        fraud=str(cell(row, "fraud")).lower() in {"1", "true", "yes", "fraud", "y"},
                    )
                )
            else:
                flows.append(
                    MoneyFlow(
                        id=new_id("flow"),
                        source_id=pid,
                        target_id=hub,
                        amount_cents=abs(amt),
                        kind="receivable",
                        memo=str(cell(row, "memo") or "AR"),
                        fraud=str(cell(row, "fraud")).lower() in {"1", "true", "yes", "fraud", "y"},
                    )
                )
        notes.append(f"Mapped {len(parties)} {role}s")

    elif kind == "supply_chain" or ("from_party" in cols and "to_party" in cols):
        kind = "supply_chain"
        for row in rows:
            src = str(cell(row, "from_party") or "").strip()
            tgt = str(cell(row, "to_party") or "").strip()
            if not src or not tgt:
                continue
            amt = parse_money(cell(row, "balance") or cell(row, "revenue"))
            tier = int(float(cell(row, "tier") or 0) or 0)
            role = str(cell(row, "role") or "vendor").strip().lower() or "vendor"
            sid, tid = slug_id("node", src), slug_id("node", tgt)
            parties.append(Party(id=sid, name=src, role=role if role != "customer" else "origin", tier=tier, amount_cents=amt))
            parties.append(
                Party(
                    id=tid,
                    name=tgt,
                    role="customer" if "cafe" in tgt.lower() or "market" in tgt.lower() else "company" if "bakery" in tgt.lower() else "distributor",
                    tier=min(tier + 1, 6),
                    amount_cents=amt,
                )
            )
            flows.append(
                MoneyFlow(
                    id=new_id("flow"),
                    source_id=sid,
                    target_id=tid,
                    amount_cents=abs(amt),
                    kind="goods",
                    memo=str(cell(row, "memo") or "supply"),
                    fraud=str(cell(row, "fraud")).lower() in {"1", "true", "yes", "fraud", "y"},
                )
            )
        notes.append(f"Mapped supply chain with {len(flows)} links")

    elif kind == "revenue" or "revenue" in cols:
        kind = "revenue"
        for row in rows:
            name = str(cell(row, "customer") or cell(row, "account_name") or "Revenue").strip()
            amt = parse_money(cell(row, "revenue") or cell(row, "balance"))
            if not amt:
                continue
            code = str(cell(row, "account_code") or "4000")
            accounts.append(
                GLAccount(code=code, name=f"Revenue · {name}", account_type="revenue", balance_cents=abs(amt))
            )
            pid = slug_id("cust", name)
            parties.append(Party(id=pid, name=name, role="customer", tier=6, amount_cents=abs(amt)))
            flows.append(
                MoneyFlow(
                    id=new_id("flow"),
                    source_id=pid,
                    target_id="maya-bakery",
                    amount_cents=abs(amt),
                    kind="revenue",
                    memo=str(cell(row, "memo") or "revenue"),
                )
            )
        notes.append(f"Mapped {len(flows)} revenue lines")

    else:
        # Generic: any amount + name columns become parties/flows
        for row in rows:
            name = str(cell(row, "vendor") or cell(row, "customer") or cell(row, "account_name") or "").strip()
            amt = parse_money(cell(row, "balance") or cell(row, "revenue"))
            if name and amt:
                pid = slug_id("party", name)
                parties.append(Party(id=pid, name=name, role="vendor", amount_cents=abs(amt)))
                flows.append(
                    MoneyFlow(
                        id=new_id("flow"),
                        source_id="maya-bakery",
                        target_id=pid,
                        amount_cents=abs(amt),
                        kind="payable",
                        memo=str(cell(row, "memo") or filename),
                    )
                )
        notes.append(f"Generic map: {len(parties)} parties")

    return {
        "kind": kind,
        "headers": headers,
        "mapped_columns": cols,
        "row_count": len(rows),
        "accounts": [a.to_dict() for a in accounts],
        "parties": [p.to_dict() for p in parties],
        "flows": [f.to_dict() for f in flows],
        "journals": [j.to_dict() for j in journals],
        "notes": notes,
        "preview_rows": [
            {headers[i]: (row[i] if i < len(row) else "") for i in range(len(headers))}
            for row in rows[:8]
        ],
    }


def rows_from_csv_bytes(data: bytes) -> tuple[list[str], list[list[Any]]]:
    text = data.decode("utf-8-sig", errors="replace")
    # Try sniffer
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",\t|;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [r for r in reader if any(str(c).strip() for c in r)]
    if not rows:
        return [], []
    return [str(h) for h in rows[0]], rows[1:]


def rows_from_xlsx_bytes(data: bytes) -> tuple[list[str], list[list[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("openpyxl required for Excel uploads — pip install openpyxl") from e
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []
    headers = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(header_row)]
    rows: list[list[Any]] = []
    for r in rows_iter:
        if r is None or not any(c is not None and str(c).strip() for c in r):
            continue
        rows.append(list(r))
    return headers, rows


def rows_from_scanned_text(text: str) -> tuple[list[str], list[list[Any]]]:
    """
    Lightweight 'OCR' path: accept pasted / extracted scan text.
    Looks for lines like: ACCOUNT 1000 Cash 18400
    or CSV-ish fragments.
    """
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if not lines:
        return [], []
    # If first line looks like headers
    if "," in lines[0] or "\t" in lines[0]:
        return rows_from_csv_bytes("\n".join(lines).encode())
    headers = ["account_code", "account_name", "balance"]
    rows: list[list[Any]] = []
    for ln in lines:
        m = re.match(
            r"(?:acct|account|gl)?\s*#?(\d{3,6})\s+(.+?)\s+\$?([\d,\.\-()]+)\s*$",
            ln,
            re.I,
        )
        if m:
            rows.append([m.group(1), m.group(2).strip(), m.group(3)])
            continue
        parts = re.split(r"\s{2,}|\t", ln)
        if len(parts) >= 3 and re.search(r"\d", parts[0]):
            rows.append([parts[0], parts[1], parts[-1]])
    return headers, rows